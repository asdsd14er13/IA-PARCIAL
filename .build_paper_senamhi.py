import json, base64, re, tempfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\Robert Herrera\escritorio\Desktop\IA PARCIAL')
NB=json.loads((ROOT/'PARCIAL IA SENAHMI.ipynb').read_text(encoding='utf-8'))
ASSET=Path(tempfile.mkdtemp(prefix='senamhi-paper-'))

def extract(cell,n,name):
    images=[]
    for o in NB['cells'][cell].get('outputs',[]):
        d=o.get('data',{}).get('image/png')
        if d: images.append(''.join(d) if isinstance(d,list) else d)
    p=ASSET/name; p.write_bytes(base64.b64decode(images[n])); return p
quality=extract(7,0,'quality.png')
station=extract(13,0,'station.png')
correlation=extract(13,1,'correlation.png')
monthly=extract(13,2,'monthly.png')
models=extract(19,0,'models.png')
importance=extract(31,0,'importance.png')

doc=Document(); sec=doc.sections[0]
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.5)
styles=doc.styles
for sn in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    s=styles[sn]; s.font.name='Times New Roman'; s._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
styles['Normal'].font.size=Pt(12); styles['Normal'].paragraph_format.line_spacing=2; styles['Normal'].paragraph_format.space_after=Pt(0); styles['Normal'].paragraph_format.first_line_indent=Inches(.5)
for sn in ['Title','Heading 1','Heading 2','Heading 3']:
    styles[sn].font.size=Pt(12); styles[sn].font.color.rgb=None; styles[sn].paragraph_format.keep_with_next=True; styles[sn].paragraph_format.space_before=Pt(10); styles[sn].paragraph_format.space_after=Pt(0); styles[sn].paragraph_format.line_spacing=2
styles['Title'].font.bold=True; styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Heading 1'].font.bold=True; styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Heading 2'].font.bold=True; styles['Heading 2'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); h._p.append(fld)

def font(run,size=12,bold=None,italic=None):
    run.font.name='Times New Roman'; run.font.size=Pt(size); run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Times New Roman'); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Times New Roman')
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    return run
def p(text='',indent=True,align=None,size=12,bold=False,italic=False):
    x=doc.add_paragraph(); x.paragraph_format.first_line_indent=Inches(.5) if indent else Inches(0); x.paragraph_format.line_spacing=2; x.paragraph_format.space_after=Pt(0)
    if align is not None: x.alignment=align
    if text: font(x.add_run(text),size,bold,italic)
    return x
def head(text,level=1):
    x=doc.add_paragraph(text,style=f'Heading {level}')
    for r in x.runs: font(r,12,bold=True)
    return x
def bullet(text):
    x=doc.add_paragraph(style='List Bullet'); x.paragraph_format.left_indent=Inches(.5); x.paragraph_format.first_line_indent=Inches(-.25); x.paragraph_format.line_spacing=2; x.paragraph_format.space_after=Pt(0); font(x.add_run(text)); return x
def caption(n,title,kind='Figura'):
    x=doc.add_paragraph(); x.paragraph_format.first_line_indent=Inches(0); x.paragraph_format.line_spacing=1; x.paragraph_format.space_before=Pt(8); x.paragraph_format.space_after=Pt(1); x.paragraph_format.keep_with_next=True; font(x.add_run(f'{kind} {n}'),10,bold=True)
    y=doc.add_paragraph(); y.paragraph_format.first_line_indent=Inches(0); y.paragraph_format.line_spacing=1; y.paragraph_format.space_after=Pt(3); y.paragraph_format.keep_with_next=True; font(y.add_run(title),10,italic=True)
def table(n,title,headers,rows,widths=None,size=8.5):
    caption(n,title,'Tabla'); t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    rp=t.rows[0]._tr.get_or_add_trPr(); rh=OxmlElement('w:tblHeader'); rh.set(qn('w:val'),'true'); rp.append(rh)
    for j,v in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=''; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'D9E2F3'); c._tc.get_or_add_tcPr().append(sh)
        z=c.paragraphs[0]; z.paragraph_format.first_line_indent=Inches(0); z.paragraph_format.line_spacing=1; z.paragraph_format.space_after=Pt(0); font(z.add_run(str(v)),size,bold=True)
    for row in rows:
        cells=t.add_row().cells; pr=t.rows[-1]._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        for j,v in enumerate(row):
            c=cells[j]; c.text=''; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; z=c.paragraphs[0]; z.paragraph_format.first_line_indent=Inches(0); z.paragraph_format.line_spacing=1; z.paragraph_format.space_after=Pt(0); font(z.add_run(str(v)),size)
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths): row.cells[j].width=Inches(w)
    p('',indent=False,size=8)
    return t
def figure(n,title,img,note,width=6.25):
    caption(n,title)
    x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.first_line_indent=Inches(0); x.paragraph_format.line_spacing=1; x.paragraph_format.keep_with_next=True
    pic=x.add_run().add_picture(str(img),width=Inches(width))
    try: pic._inline.docPr.set('descr',title); pic._inline.docPr.set('title',title)
    except Exception: pass
    y=doc.add_paragraph(); y.paragraph_format.first_line_indent=Inches(0); y.paragraph_format.line_spacing=1; y.paragraph_format.space_after=Pt(5); font(y.add_run('Nota. '),9,italic=True); font(y.add_run(note),9)

# APA 7 student title page, using the author and course details supplied by the student.
for _ in range(3): p('',indent=False)
p('Comparación de modelos de regresión para estimar PM2.5 con datos horarios abiertos de SENAMHI en Lima Metropolitana (2015–2024)',False,WD_ALIGN_PARAGRAPH.CENTER,12,True)
for _ in range(2): p('',indent=False)
for v in ['Robert Eloy Herrera Ccari','Universidad Nacional Federico Villarreal','Ingeniería Informática','Curso de Inteligencia Artificial','Ciro Rodríguez Rodríguez','27 de septiembre de 2026']: p(v,False,WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

abstract=('Este trabajo compara modelos de aprendizaje automático para estimar la concentración horaria de PM2.5 en siete estaciones de Lima Metropolitana. Se utilizaron datos abiertos del Servicio Nacional de Meteorología e Hidrología del Perú (SENAMHI), con 577 794 registros entre enero de 2015 y mayo de 2024. Luego de convertir fechas, eliminar duplicados exactos y consolidar registros repetidos por estación y hora, quedaron 577 710 filas. Para el experimento principal se seleccionaron 209 102 observaciones con valores disponibles de PM2.5, PM10 y NO2. La tarea se definió como estimación contemporánea: el modelo usa PM10 y NO2 medidos en la misma hora, además de la estación y variables temporales; por ello, no constituye un pronóstico de horas futuras. Se compararon un baseline de la media, regresión lineal, K vecinos más cercanos, árbol de decisión, Random Forest y un Voting que promedia regresión lineal y Random Forest. Los datos se dividieron cronológicamente, con entrenamiento hasta diciembre de 2022 y prueba desde enero de 2023 hasta mayo de 2024. Debido al costo de ejecución se usaron muestras reproducibles de 30 000 registros de entrenamiento y 15 000 de prueba. Random Forest obtuvo el menor error en la comparación principal (MAE = 5.165; RMSE = 8.144; R² = .499), mientras que Voting obtuvo RMSE = 8.300. El Random Forest superó al baseline, pero el valor de R² muestra que aún queda variabilidad sin explicar. Las combinaciones adicionales ensayadas no demostraron una mejora bajo sus protocolos actuales. El modelo guardado puede reutilizarse con registros compatibles, pero no debe interpretarse como una alerta oficial ni como una predicción futura.')
abstract=('Este trabajo compara modelos de aprendizaje automático para estimar la concentración horaria de PM2.5 en siete estaciones de Lima Metropolitana. Se utilizaron datos abiertos del Servicio Nacional de Meteorología e Hidrología del Perú (SENAMHI), con 577 794 registros entre enero de 2015 y mayo de 2024. Luego de convertir fechas, eliminar duplicados exactos y consolidar registros repetidos por estación y hora, quedaron 577 710 filas. Para el experimento principal se seleccionaron 209 102 observaciones con valores disponibles de PM2.5, PM10 y NO2. La tarea se definió como estimación contemporánea: el modelo usa PM10 y NO2 medidos en la misma hora, además de la estación y variables temporales; por ello, no constituye un pronóstico de horas futuras. Se compararon un baseline de la media, regresión lineal, K vecinos más cercanos, árbol de decisión, Random Forest y un Voting que promedia regresión lineal y Random Forest. Los datos se dividieron cronológicamente, con entrenamiento hasta diciembre de 2022 y prueba desde enero de 2023 hasta mayo de 2024. Debido al costo de ejecución se usaron muestras reproducibles de 30 000 registros de entrenamiento y 15 000 de prueba. Random Forest obtuvo el menor error en la comparación principal (MAE = 5.165; RMSE = 8.144; R² = .499), mientras que Voting obtuvo RMSE = 8.300. El Random Forest superó al baseline, pero el valor de R² muestra que aún queda variabilidad sin explicar. Las combinaciones adicionales ensayadas no demostraron una mejora bajo sus protocolos actuales. El modelo guardado puede reutilizarse con registros compatibles, pero no debe interpretarse como una alerta oficial ni como una predicción futura.')
wc=len(re.findall(r'\b[\wÁÉÍÓÚáéíóúñÑ²]+\b',abstract)); assert 200<=wc<=300,wc
head('1. Título',1); p('Comparación de modelos de regresión para estimar PM2.5 con datos horarios abiertos de SENAMHI en Lima Metropolitana (2015–2024).')
head('2. Resumen',1); p(abstract,False)
head('3. Palabras clave',1)
x=doc.add_paragraph(); x.paragraph_format.first_line_indent=Inches(.5); x.paragraph_format.line_spacing=2; font(x.add_run('Calidad del aire; PM2.5; aprendizaje automático; regresión; Random Forest; Lima Metropolitana.'))

head('4. Introducción',1)
p('La contaminación del aire es un tema relevante para la salud y la gestión ambiental. Entre los contaminantes monitoreados se encuentra el material particulado fino PM2.5, definido por su diámetro aerodinámico de hasta 2.5 micrómetros. La Organización Mundial de la Salud (OMS, 2021) publicó recomendaciones para PM2.5 y otros contaminantes; estas guías sanitarias no deben confundirse con los estándares legales peruanos. En el Perú, el Decreto Supremo N.° 003-2017-MINAM fija los Estándares de Calidad Ambiental (ECA) para aire. Para PM2.5, el ECA diario es de 50 µg/m³ (Ministerio del Ambiente, 2017; SENAMHI, 2024b).')
p('Los boletines oficiales muestran cambios por estación y periodo. En enero de 2024, SENAMHI reportó excedencias del ECA diario en SMP, CRS y CRB, con un máximo de 86.46 µg/m³ en Carabayllo. En agosto de 2026, el boletín mensual reportó un máximo de 63.9 µg/m³ en Pariachi. Son episodios localizados, no valores representativos de toda Lima ni comparables directamente con estimaciones horarias (SENAMHI, 2024b, 2026).')
p('En Lima ya se han estudiado diferentes formas de analizar contaminantes. Montalvo et al. (2022) describieron un sistema local con sensores de bajo costo y modelos de IA que genera mapas y pronósticos de seis horas. Rojas et al. (2022) y Sánchez-Ccoyllo y Alonso (2024) estudiaron concentraciones de PM mediante simulación atmosférica y modelos de transporte químico. Solis Teran et al. (2025) compararon modelos estadísticos y redes neuronales para pronosticar PM10, mientras que Gavidia et al. (2026) clasificaron categorías de calidad del aire con registros horarios de SENAMHI.')
p('Este trabajo pregunta qué regresor estima con menor error el PM2.5 de la misma hora si se conocen PM10, NO2, estación y variables temporales. Se comparan varios modelos y un baseline mediante una separación cronológica. El aporte es documentar una comparación reproducible con datos locales abiertos y guardar un pipeline; no se pretende sustituir instrumentos oficiales ni emitir alertas.')

head('5. Situación problemática',1)
p('El recurso de SENAMHI contiene registros horarios validados de PM10, PM2.5 y NO2, pero presenta faltantes: 36.38 % en PM2.5, 34.06 % en PM10 y 48.22 % en NO2. Además, el análisis se limita a las siete estaciones y al periodo del archivo. Los faltantes describen disponibilidad, pero no prueban por sí solos una falla de monitoreo.')
p('En enero de 2024 se registraron concentraciones diarias de PM2.5 superiores al ECA de 50 µg/m³ en SMP, CRS y CRB; Carabayllo alcanzó 86.46 µg/m³. En agosto de 2026, SENAMHI describió una excedencia en Pariachi con 63.9 µg/m³. Son promedios diarios y periodos distintos a la prueba horaria del notebook; se citan como evidencia de episodios localizados, no como resultado del modelo (SENAMHI, 2024b, 2026).')
p('La pregunta práctica es si las otras mediciones horarias permiten estimar PM2.5 para una observación cuando su etiqueta está disponible en el histórico. El experimento no evalúa alertas futuras ni una recuperación operativa de lecturas faltantes en una red en producción.')

head('6. Identificación y formulación del problema',1)
head('Problema general',2); p('No se ha determinado, para este conjunto y bajo una prueba temporal reproducible, qué modelo estima PM2.5 con menor error a partir de estación, tiempo, PM10 y NO2 de la misma hora.')
head('Causas y factores relacionados',2); p('Los registros tienen campos vacíos; la disponibilidad varía por estación y fecha; y las relaciones entre contaminantes pueden no ser lineales. El análisis describe estos factores, pero no prueba que sean causas de la concentración.')
head('Consecuencias y necesidad',2); p('Sin una comparación local resulta difícil seleccionar una alternativa para este ejercicio. Se necesita contrastar varios modelos con un test temporal común y explicar los límites de la tarea.')
head('Evidencias',2); p('La evidencia incluye los conteos de faltantes y duplicados del CSV, las siete estaciones, las 209 102 filas completas de modelado y los boletines oficiales que muestran variabilidad de PM2.5.')

head('7. Pregunta de investigación',1)
p('¿Qué desempeño presentan distintos modelos de regresión para estimar las concentraciones horarias de PM2.5 usando estación, variables temporales y mediciones de PM10 y NO2 de la misma hora en los datos abiertos de monitoreo de Lima Metropolitana?')

head('8. Objetivos',1)
head('Objetivo general',2); p('Comparar modelos de aprendizaje automático para estimar PM2.5 en los registros de monitoreo de Lima Metropolitana, utilizando datos abiertos de SENAMHI y una evaluación temporal, con el fin de identificar el modelo con menor error en el periodo de prueba.')
head('Objetivos específicos',2)
for s in ['OE1. Describir la estructura, cobertura y calidad de los datos de SENAMHI.','OE2. Preparar los registros y documentar el tratamiento de fechas, duplicados, faltantes y variables.','OE3. Entrenar y comparar un baseline, modelos de regresión y combinaciones mediante una partición temporal.','OE4. Evaluar los resultados, identificar el modelo con menor error y guardar su pipeline para datos compatibles.']: bullet(s)
table(1,'Relación entre pregunta, objetivos, experimentos y evidencia',['Problema / pregunta','Objetivo','Método','Evidencia'],[
('Estimar PM2.5 con datos reales.','OE1','Describir la fuente, periodo, estaciones y faltantes.','Ficha y tabla de calidad.'),
('La base tiene vacíos y horas repetidas.','OE2','Convertir fechas y consolidar estación-hora con mediana.','Conteos antes/después.'),
('No se conoce qué modelo minimiza el error.','OE3','Comparar baseline, regresores y Voting.','Métricas del test temporal.'),
('Se necesita reutilizar el modelo elegido.','OE4','Evaluar y guardar el mejor pipeline.','Tabla comparativa y ficha joblib.')], [1.8,.7,2.3,1.65],8.2)

head('9. Justificación',1)
p('La justificación práctica es académica: una comparación local permite observar si los datos abiertos contienen información útil para estimar PM2.5 en la misma hora. La justificación tecnológica es evaluar aprendizaje automático frente a un baseline, sin asumir que el modelo más complejo será superior. Las revisiones de Yin (2025) y Karmoude et al. (2025) muestran que el análisis de PM2.5 incluye estimación, predicción y distintos tipos de modelos; el desempeño depende de la tarea y los datos.')
p('La justificación científica es acotada: se deja evidencia de una comparación reproducible con datos de siete estaciones de Lima y una prueba posterior en el tiempo. La justificación metodológica es organizar el proceso en comprensión, preparación, modelado y evaluación, siguiendo el esquema general CRISP-DM (Wirth & Hipp, 2000). El resultado no es una herramienta oficial de vigilancia.')

head('10. Antecedentes y estado del arte',1)
p('Se seleccionaron estudios que ubican el problema local, comparan tareas y muestran que un ensamble puede o no mejorar según el caso. No se comparan sus métricas numéricamente cuando cambian la variable objetivo, ciudad, entradas o validación.')
table(2,'Estudios relacionados de monitoreo, regresión y ensambles',['Estudio','Método y hallazgo reportado','Relación y diferencia con este proyecto'],[
('Montalvo et al. (2022), Lima','Sensores de bajo costo, IA y mapas; evalúa predicción espacial/temporal y pronósticos de 6 h.','Antecedente local; tarea diferente a estimación contemporánea.'),
('Rojas et al. (2022), Lima','CCATT-BRAMS y datos SENAMHI para simular PM y analizar meteorología.','Contexto local de PM; método físico, no mismos regresores.'),
('Sánchez-Ccoyllo y Alonso (2024), Lima-Callao','BRAMS 5.2 y emisiones vehiculares para mejorar simulación de PM10/PM2.5.','Antecedente peruano de modelado, distinto a ML supervisado.'),
('Yang (2026), Londres','Random Forest para concentraciones diarias de PM2.5 y AQI; el artículo señala PM10 como predictor importante.','Comparte RF y contaminantes; su horizonte y frecuencia diarios difieren de la estimación horaria.'),
('Izabi et al. (2026), Makassar','Random Forest sobre 395 registros diarios; compara dos particiones y trata faltantes con interpolación y regresión múltiple.','Apoya la comparación de RF en aire, pero es otra ciudad, muestra pequeña y estrategia de faltantes distinta.'),
('Singh et al. (2026), Taiwán','Voting ponderado de cuatro modelos de boosting para AQI; usa 74 estaciones y evalúa con cortes temporal y aleatorio, además de SHAP.','Ejemplo de ensamble especializado para AQI; objetivo, algoritmos y tamaño de datos no equivalen a PM2.5 horario en Lima.'),
('Tian et al. (2024), Macao','Stacking LSTM + XGBoost con LightGBM meta; superó métodos comparados en su conjunto.','Ensamble exitoso en otro contexto; no garantiza mejora aquí.'),
('Solis Teran et al. (2025), Lima','Series de tiempo y LSTM para PM10; ARIMA y SSA tuvieron menor error medio que LSTM.','Ejemplo local de modelo tradicional competitivo; objetivo distinto.'),
('Gavidia et al. (2026), Lima','Gaussian Naïve Bayes para categorías con 768 185 registros SENAMHI y división temporal.','Antecedente local de clasificación; no predice valor numérico.'),
('Alvarez et al. (2025), Quito','Modelos con representaciones satelitales AlphaEarth para concentraciones anuales y mapas; compara varios regresores.','Muestra una vía espacial distinta; no cuenta con los mismos datos horarios ni predictores.'),
('Peinado et al. (2025), revisión sistemática','Sintetiza 412 estudios y compara familias estadísticas y de ML para calidad del aire.','Sustenta la selección contextual de modelos, no predice el desempeño de nuestro archivo.'),
('Karmoude et al. (2025); Yin (2025)','Revisiones de ML y análisis de PM2.5.','Ubican métodos, tareas y límites de generalización.')],[1.3,2.75,2.4],7.8)
p('Los estudios locales muestran que se han usado sensores económicos, simulación atmosférica, redes neuronales y clasificación con datos de Lima. Su comparación directa con este proyecto sería incorrecta si se ignoraran las diferencias entre pronosticar una hora futura, estimar un valor de la misma hora, simular químicamente la atmósfera o asignar una categoría de AQI. Aquí el resultado buscado es un número de PM2.5 contemporáneo para un registro con PM10 y NO2 disponibles.')
p('Yang (2026) y los trabajos de Izabi et al. (2026) y Muralikrishnan et al. (2026) respaldan el uso de modelos de árboles en otros conjuntos de aire, pero no demuestran que Random Forest sea universalmente el mejor. Izabi et al. trabajaron con 395 observaciones diarias de Makassar y combinaron interpolación lineal con regresión múltiple para tratar faltantes; este notebook, en cambio, conserva solo filas completas en PM2.5, PM10 y NO2 para el experimento principal. El artículo de Singh et al. (2026) probó un Voting ponderado para AQI de Taiwán, con cuatro modelos de boosting, 74 estaciones y una muestra mucho mayor. Su resultado es evidencia de que ciertos ensambles pueden funcionar en otra configuración, no una razón para afirmar que el Voting lineal–bosque de Lima deba mejorar.')
p('La revisión de Peinado et al. (2025) resume 412 trabajos de calidad del aire y describe que el rendimiento depende de disponibilidad de datos, resolución temporal o espacial y objetivo. Alvarez et al. (2025) predijeron concentraciones anuales en Quito usando características satelitales y compararon varios algoritmos; ese diseño geoespacial difiere de nuestras mediciones estación-hora. Como apoyo conceptual, Kovač et al. (2024) ofrecen una guía aplicada para regresión con Python, exploración de correlación, preprocesamiento y evaluación; se usa por su explicación práctica, aunque sus ejemplos pertenecen a psicología y no a contaminación.')

head('11. Marco conceptual',1)
head('PM2.5, PM10 y NO2',2); p('PM2.5 identifica partículas finas de hasta 2.5 µm; PM10, partículas de hasta 10 µm; NO2 es dióxido de nitrógeno. El CSV registra estas variables por estación y hora. Los boletines de SENAMHI reportan concentraciones en µg/m³, pero el diccionario local no declara explícitamente la unidad en las columnas, por lo que se recomienda confirmarla.')
head('Regresión y modelos',2); p('La regresión estima un número. El baseline predice la media; la regresión lineal representa relaciones lineales; KNN estima a partir de vecinos; un árbol genera reglas; Random Forest promedia árboles para representar patrones no lineales (Breiman, 2001). Voting promedia las predicciones de modelos diferentes. Un ensamble ayuda si reduce errores; añadir modelos no garantiza una mejora.')
head('Métricas e interpretación',2); p('MAE es el error absoluto promedio; RMSE penaliza más los errores grandes; R² compara el error con la variación alrededor de la media y puede ser negativo. La importancia por permutación mide cuánto cambia el score al mezclar una variable, pero no demuestra causalidad.')
head('Estimación contemporánea y pronóstico',2); p('PM10 y NO2 se miden a la misma hora que el PM2.5 objetivo. Por ello, el modelo estima PM2.5 contemporáneo cuando conoce esos valores; no anticipa horas futuras. Un pronóstico requeriría variables disponibles antes del instante objetivo y otra validación.')

head('12. Metodología de investigación',1)
p('El enfoque es cuantitativo y aplicado: se analizan mediciones numéricas reales y se construye un artefacto de aprendizaje automático. El diseño es comparativo y experimental porque se entrenan modelos alternativos y se evalúan en el mismo periodo temporal para el análisis principal. La unidad de análisis es un registro consolidado estación-hora. No se interviene sobre estaciones ni se recolectan datos personales.')
p('También se puede describir como un estudio retrospectivo de datos secundarios: el equipo no instaló sensores ni generó las observaciones, sino que tomó el archivo público ya recogido por SENAMHI. La comparación se concentra en un resultado cuantitativo (PM2.5) y en un conjunto delimitado de predictores, y no pretende explicar causalmente por qué sube o baja la contaminación. Esta precisión evita confundir un ejercicio de modelado predictivo con un estudio experimental de emisiones o salud.')
p('El trabajo construye y evalúa un artefacto de software: un pipeline que organiza el preprocesamiento y el estimador para reutilizarse con entradas compatibles. Peffers et al. (2012) analizan cómo evaluar artefactos de investigación de diseño; se toma esa idea de evaluar el artefacto frente a un criterio explícito, el error en el test temporal. Sin embargo, este parcial no desarrolla todas las fases formales de Design Science Research ni prueba el artefacto con usuarios en una institución, por lo que se presenta como un proyecto aplicado con evaluación comparativa, no como una implementación organizacional completa.')
head('13. Metodología de desarrollo de IA',1)
p('El notebook sigue etapas didácticas de CRISP-DM: comprensión del problema, comprensión y preparación de datos, modelado, evaluación y guardado del pipeline. El guardado local no equivale a despliegue productivo. CRISP-DM organiza de manera iterativa proyectos de minería de datos (Wirth & Hipp, 2000).')
p('La guía de ciclo de vida de Sarker (2021) presenta las familias principales de aprendizaje automático y sus usos generales. Kosmas et al. (2026) y Fang et al. (2023) describen actividades que acompañan el ciclo de vida, desde datos y entrenamiento hasta evaluación, seguimiento y operación. Estas fuentes ayudan a explicar por qué no basta con llamar a `.fit()`: también se define el objetivo, se revisa la calidad, se prepara el pipeline, se compara el resultado y se documenta cómo se guarda. El alcance de este curso termina en un artefacto local probado con el mismo esquema de datos; no incluye monitoreo continuo, control de deriva ni infraestructura MLOps.')
p('La secuencia concreta fue iterativa. Primero se definió que PM2.5 era un objetivo numérico. Luego se comprobó cómo estaban fechados los registros y qué proporción de cada contaminante estaba vacía. Con esa evidencia se eligió una evaluación temporal y se construyeron variables cíclicas. Después se compararon modelos individuales y combinaciones. Finalmente se registró el modelo que obtuvo el menor RMSE entre los regresores principales y se dejó una función para cargar el pipeline con `joblib`. En cada decisión se mantuvo visible la diferencia entre predicción futura y estimación con mediciones simultáneas.')
head('14. Fase 1: comprensión del problema',1)
p('El prototipo recibe estación, fecha, hora, PM10 y NO2 y devuelve un PM2.5 estimado de la misma observación. El criterio principal es menor RMSE, acompañado de MAE y R². La restricción esencial es que no genera alertas oficiales ni predice horas futuras.')

head('15. Fase 2: comprensión de los datos',1)
p('La fuente es el conjunto “Monitoreo de los contaminantes del aire en Lima Metropolitana”, de SENAMHI, disponible en la Plataforma Nacional de Datos Abiertos. Los metadatos indican mediciones horarias validadas, formato CSV, actualización mensual, licencia Open Data Commons Attribution y cobertura 2015–2024. El archivo estudiado contiene 577 794 filas, 15 columnas y siete estaciones hasta el 31 de mayo de 2024. El boletín de 2026 describe una red actual de diez estaciones, que no debe confundirse con las siete del CSV analizado (SENAMHI, 2024a, 2026).')
table(3,'Ficha resumida del conjunto de datos',['Característica','Descripción'],[
('Entidad/fuente','SENAMHI, Plataforma Nacional de Datos Abiertos.'),('Periodo','2015-01-01 a 2024-05-31.'),('Tamaño','577 794 filas, 15 columnas.'),('Estaciones','Siete estaciones en Lima Metropolitana.'),('Unidad','Medición estación-hora; claves repetidas se consolidan.'),('Objetivo','PM2_5 numérico.'),('Entradas','ESTACION, PM10, NO2, hora, día semana, mes y variables cíclicas.'),('Licencia','Open Data Commons Attribution, según metadatos.'),('Unidad física','El diccionario local no la señala; boletines reportan µg/m³. Confirmar con SENAMHI.')],[1.75,4.7],8.8)

head('16. Ingeniería y preparación de datos',1)
p('La preparación comenzó con una copia del CSV original, de 577 794 filas y 15 columnas. Se retiraron espacios de los nombres de columna y se uniformizó el texto de ESTACION con eliminación de espacios laterales y mayúsculas. Los campos de contaminantes y coordenadas se convirtieron a formato numérico; FECHA y HORA se reunieron en una marca temporal para ordenar el registro y derivar variables de calendario. No se descartó ninguna fila por una fecha inválida, y no se encontraron duplicados exactos.')
p('El control por clave ESTACION–timestamp detectó 42 claves repetidas, asociadas a 84 filas señaladas como repetidas en la salida del notebook. Para no contar más de una vez la misma hora de una estación, se agruparon los valores de contaminantes con la mediana y se conservó una sola fila por clave. La mediana reduce el efecto de un valor extremo dentro de una hora repetida, pero no garantiza que el dato resultante represente una lectura físicamente correcta; por eso se reporta como regla de consolidación, no como validación instrumental. Tras este paso quedaron 577 710 filas.')
p('La exploración contabilizó los valores ausentes: 196 812 en PM10 (34.06 %), 210 203 en PM2.5 (36.38 %) y 278 602 en NO2 (48.22 %). El cuaderno muestra cómo se distribuyen por año y contaminante para que se vea que la disponibilidad no es uniforme. En el análisis principal no se imputaron PM2.5, PM10 o NO2 cuando faltaban: se conservaron únicamente registros con objetivo y los dos contaminantes de entrada disponibles. Este filtro dejó 209 102 observaciones. Por lo tanto, las métricas describen ese subconjunto completo y no representan automáticamente los registros excluidos.')
p('La mediana de PM2.5 por mes agrega las estaciones disponibles. Es útil para notar variación temporal en la serie, pero no corrige cambios en cobertura: si en un mes hay menos estaciones u horas válidas, la composición de la mediana puede cambiar. El boxplot por estación resume diferencias entre distribuciones y oculta los puntos extremos solo en la visualización; esos puntos no se retiraron del entrenamiento por esa razón. Los tres contaminantes tuvieron cero valores negativos en las verificaciones del notebook, lo cual es un control básico y no equivale a un control de calibración de sensores.')
p('El artículo de Mu et al. (2026) revisa el preprocesamiento en modelos de energía de edificios y organiza el trabajo en análisis de datos, preparación e ingeniería de variables. Se cita como marco general de flujo y se aclara que su evidencia se refiere a edificios, no a mediciones de aire. En cambio, Cicek y Aydin (2026) estudian específicamente imputación y pronóstico de contaminantes en estaciones de Turquía y comparan varios métodos de reconstrucción bajo faltantes simulados. Ese artículo ayuda a explicar que imputar es una decisión metodológica con consecuencias, pero nuestro notebook no implementó su técnica ni un método de imputación para los campos requeridos en el experimento principal.')
p('Para representar patrones cíclicos se añadieron hora, día de semana y mes como variables de calendario. Además, hora se codificó con seno y coseno usando un periodo de 24 horas, y día del año con seno y coseno usando 365.25 días. Esta representación evita que el final y el inicio del ciclo queden muy separados numéricamente: por ejemplo, las 23:00 y las 00:00 son horas vecinas en el reloj. Alsabagh et al. (2025) ensayan codificaciones cíclicas en un problema horario de contaminación en Beijing; sus modelos profundos y resultados no se transfieren a Lima, pero el uso de una codificación circular es pertinente para describir esta decisión de ingeniería.')
table(4,'Calidad de datos y preparación',['Control','Resultado','Lectura'],[
('Fechas inválidas','0','No hubo descartes por fecha.'),('Duplicados exactos','0','No se encontraron filas idénticas.'),('Claves repetidas','84 filas / 42 claves','Consolidadas con mediana.'),('Filas limpias','577 710','Después de consolidar.'),('PM10 faltante','196 812 (34.06%)','No se imputó para el caso completo principal.'),('PM2.5 faltante','210 203 (36.38%)','No se usa como etiqueta si falta.'),('NO2 faltante','278 602 (48.22%)','Se excluye del caso completo.'),('Filas completas','209 102','166 568 antes de 2023; 42 534 desde 2023.')],[1.45,1.85,3.15],8.3)
figure(1,'Valores faltantes y mediciones disponibles por año',quality,'La disponibilidad cambia por contaminante y periodo; la figura no identifica por sí sola la causa de los faltantes.')
figure(2,'Mediciones válidas por año y distribución de PM2.5 por estación',station,'El boxplot oculta puntos extremos para facilitar lectura; esos valores no fueron eliminados del modelado.')
figure(3,'Correlación descriptiva entre PM10, PM2.5 y NO2',correlation,'La matriz resume asociación lineal entre contaminantes en las filas disponibles; correlación no implica causalidad ni demuestra que una variable cause otra.')
figure(4,'Mediana mensual de PM2.5 con estaciones agregadas',monthly,'Cada punto es la mediana mensual de las mediciones válidas agrupadas; la composición de estaciones y horas disponibles puede variar entre meses.')
p('Las Figuras 3 y 4 cumplen funciones descriptivas distintas. El mapa de correlación permite revisar si los contaminantes tienden a variar juntos en el conjunto observado, mientras que la serie mensual resume la mediana en el tiempo. Ninguna se utilizó para seleccionar el periodo de prueba ni para cambiar etiquetas después de observarlas. Tampoco se debe interpretar la serie agregada como el comportamiento de cada estación: para eso se necesitarían gráficos y métricas separados por estación.')

head('17. Diseño de la solución',1)
p('La solución se organiza como una tubería que recibe un registro con estación, fecha, hora, PM10 y NO2. El cuaderno transforma la fecha y hora, deriva las variables de calendario, aplica la codificación de estación y ejecuta el Random Forest seleccionado. La salida es una estimación de PM2.5 para esa misma observación. El diseño mantiene juntos los pasos de preparación y el estimador, de modo que al reutilizarlo no sea necesario recrear manualmente las reglas de escalado y codificación.')
p('La elección se hace después de comparar modelos y no por anticipado: primero el pipeline evalúa baseline, regresión lineal, KNN, árbol, Random Forest y Voting; luego el archivo de uso guarda el Random Forest individual porque obtuvo el menor RMSE en la comparación principal. Los análisis de clasificación y K-Means responden preguntas auxiliares distintas y no se conectan a la función que estima PM2.5. Esta separación evita presentar una etiqueta o un grupo como si fuera la predicción de la concentración.')
caption(5,'Flujo del artefacto de estimación')
p('Datos abiertos SENAMHI  →  limpieza y consolidación  →  variables de estación y hora  →  pipeline + Random Forest  →  estimación de PM2.5 para la misma hora',False,WD_ALIGN_PARAGRAPH.CENTER,10)
p('Nota. El modelo utiliza PM10 y NO2 medidos en el mismo instante; elaboración propia a partir del notebook.',False,size=9)
head('Requisitos principales',2)
bullet('Funcional: recibir estación, fecha/hora, PM10 y NO2; preparar variables y devolver PM2.5 estimado.')
bullet('No funcional: conservar preprocesamiento y modelo juntos, registrar versiones y semilla.')
bullet('Restricción: no usar como predicción futura si no se conocen PM10 y NO2 del instante objetivo.')

head('18. Desarrollo e implementación',1)
head('Variables y pipeline',2)
p('Las diez variables de entrada son ESTACION, PM10, NO2, hora, día de semana, mes, hora_sin, hora_cos, dia_sin y dia_cos. PM2_5 es la variable objetivo y no se incluye como entrada. Para variables numéricas, el preprocesamiento contiene una imputación por mediana y StandardScaler; para ESTACION, contiene imputación por categoría más frecuente y One-Hot Encoding con tolerancia a estaciones nuevas. Como la tabla principal ya filtra PM2.5, PM10 y NO2 completos, la imputación no reconstruye los casi 370 mil registros ausentes de los campos contaminantes. Su presencia permite conservar el pipeline de transformación coherente para valores que sí entren al modelo y manejar categorías desconocidas en estación.')
p('Se compararon el baseline de la media, una regresión lineal, KNN, un árbol de decisión, Random Forest y Voting. El baseline asigna a todos los registros la media de PM2.5 del entrenamiento y sirve para saber si un modelo aporta frente a una regla simple. La regresión lineal busca una combinación lineal de predictores. KNN toma vecinos similares en el espacio preparado y combina sus salidas con pesos por distancia. Un árbol divide variables en reglas sucesivas. Random Forest entrena muchos árboles sobre muestras internas y promedia sus salidas, lo que puede capturar relaciones no lineales sin depender de un único árbol (Breiman, 2001). Voting toma el promedio aritmético de las predicciones del componente lineal y del componente Random Forest. Las transformaciones y modelos se implementaron como estimadores y pipelines de scikit-learn (Pedregosa et al., 2011).')
table(5,'Configuración de los regresores de la comparación principal',['Modelo','Configuración registrada en el notebook','Función en la comparación'],[
('Baseline de la media','Promedio de PM2_5 en entrenamiento; mismo valor para cada caso de prueba.','Referencia simple; no tiene hiperparámetros de árbol o distancia.'),
('Regresión lineal','LinearRegression() con valores predeterminados de scikit-learn.','Referencia lineal multivariable.'),
('KNN regresor','7 vecinos; pesos por distancia; n_jobs = −1; distancia Minkowski por defecto (p = 2).','Estimación por observaciones cercanas.'),
('Árbol regresor','Profundidad máxima = 14; mínimo 5 casos por hoja; semilla = 42.','Reglas no lineales mediante particiones.'),
('Random Forest regresor','100 árboles; profundidad máxima = 18; mínimo 3 casos por hoja; semilla = 42; n_jobs = −1.','Modelo individual que obtuvo menor RMSE.'),
('Voting regresor','Promedio con pesos iguales de LinearRegression y RandomForest (mismos 100 árboles y límites).','Comprueba si dos predicciones complementarias reducen el error.')],[1.5,3.15,1.8],7.8)
p('Los valores “por defecto” de la tabla son los que quedaron definidos por la biblioteca en la versión guardada; el notebook fija explícitamente solo los parámetros escritos. No se realizó una búsqueda exhaustiva de hiperparámetros para cada estimador. Por ello, la comparación responde qué ocurrió con estas configuraciones concretas, no cuál es la configuración óptima posible para cada familia de modelos. Kovač et al. (2024) muestran en una guía práctica cómo un trabajo de regresión pasa por exploración, preparación y evaluación; se usa como referencia didáctica para reportar estas etapas, no como evidencia sobre contaminación atmosférica.')
head('Modelo guardado',2)
p('El archivo `modelo_pm25_senamhi.joblib` contiene un diccionario con una sección `modelos` y otra `ficha`. En `modelos` está únicamente la entrada `mejor_individual`, que apunta al pipeline del Random Forest regresor. La ficha registra las variables requeridas, corte temporal, tamaño de las muestras, semilla, versiones de pandas y scikit-learn, métricas y resultados de las combinaciones ensayadas. No contiene modelos de Voting, mezcla ponderada, stacking, clasificación ni K-Means como candidatos de inferencia. La ficha describe experimentos, pero no equivale a guardar cada modelo para predicción.')
p('Para usarlo, la función del notebook recibe estación, FECHA, HORA, PM10 y NO2, genera las características temporales y carga el pipeline. Para la estación, el codificador one-hot puede ignorar categorías no vistas durante el ajuste, aunque eso no garantiza que el modelo generalice bien a una estación nueva. La tabla de predicción devuelve PM2_5_estimado y el nombre del modelo utilizado. La unidad de esta salida debe confirmarse con el diccionario o el proveedor antes de presentarla como una concentración física en µg/m³.')

head('19. Diseño experimental',1)
p('Antes de muestrear, el subconjunto completo se separó por fecha. El periodo con objetivo y predictores presentes fue 2015-02-03–2022-12-31 para entrenamiento disponible (166 568 filas) y 2023-01-01–2024-05-31 para prueba disponible (42 534 filas). Se eligió una separación cronológica porque la pregunta se refiere a evaluar datos posteriores al entrenamiento. La preparación se incluye dentro del pipeline y se ajusta con la muestra de entrenamiento, con lo que la escala numérica y las categorías de estación no se calculan usando la prueba.')
p('Por el tiempo de cómputo, en el experimento principal se seleccionaron aleatoriamente 30 000 registros de entrenamiento y 15 000 de prueba, con semilla 42; luego se ordenaron por fecha para facilitar inspección. La semilla permite repetir la selección si el archivo y el entorno se mantienen. No convierte la muestra en un censo ni elimina la incertidumbre que produce evaluar solo 15 000 de los 42 534 registros elegibles. Las seis alternativas principales recibieron el mismo conjunto muestreado de entrenamiento y test, condición necesaria para una comparación directa dentro de ese experimento.')
p('La separación cronológica responde al principio de no usar observaciones posteriores para entrenar un modelo que se evalúa en periodos siguientes. Fauszt (2026) formula la validez temporal junto con la definición del momento de predicción y de la información que puede conocerse entonces; Mishra et al. (2026) clasifican riesgos de fuga temporal en una aplicación de predicción de compilaciones de software. Los dos estudios pertenecen a contextos distintos, así que se toman como advertencia metodológica general y no como evidencia específica de SENAMHI. Nassar (2023), en un informe breve de autor, también explica que la mezcla inadvertida entre entrenamiento, validación y prueba puede inflar métricas. En este notebook, el corte por fecha y el ajuste del preprocesamiento solo sobre train reducen esos riesgos, pero no resuelven todos los problemas de validez externa o de predictores contemporáneos.')
p('Se hicieron tres preguntas experimentales. La primera fue si los modelos individuales superaban al baseline de la media en una prueba posterior a los años de entrenamiento. La segunda fue si un promedio igual de regresión lineal y Random Forest mejoraba el RMSE de cada componente. La tercera buscó una combinación más flexible usando validación de 2022, sin elegir pesos mirando el test de 2023 en adelante. Estas preguntas deben leerse por separado porque los experimentos segundo y tercero no usaron el mismo volumen o periodo de entrenamiento que el Random Forest de la tabla principal.')
p('Para la mezcla ponderada, se entrenaron cuatro modelos base con datos anteriores al 1 de enero de 2022; se calcularon sus predicciones en una muestra de 15 000 casos de 2022 y se revisaron todos los pares de modelos. Para cada par se evaluaron 21 pesos uniformes entre 0 y 1, con paso de 0.05, minimizando RMSE de validación. La mejor configuración fue 0 % regresión lineal y 100 % Random Forest (RMSE 14.648 en validación 2022); matemáticamente, esa mezcla seleccionada equivale a usar solo el bosque. Después se midió en el periodo de prueba y su RMSE fue 9.430.')
p('Para stacking se reutilizaron cuatro modelos base ajustados antes de 2022: regresión lineal, KNN, árbol y Random Forest. Sus predicciones de 2022 alimentaron un metaestimador RidgeCV, que selecciona entre alfas 0.1, 1, 10 y 100. La prueba final fue de 2023 en adelante. Así el modelo de nivel superior aprende combinaciones a partir de un periodo de validación y se evalúa en un periodo posterior. En cambio, el Random Forest principal entrenó con datos hasta diciembre de 2022. Esa diferencia significa que el test de fechas puede coincidir, pero el tamaño y los años usados para ajustar los componentes cambian; por eso el resultado del stacking es exploratorio y no una comparación completamente controlada del efecto del ensamble.')

head('20. Evaluación',1)
p('La medida principal fue RMSE, raíz del promedio de los errores al cuadrado. Al elevar al cuadrado, los errores grandes pesan más; por eso sirve si preocupa especialmente una predicción muy alejada del valor observado. MAE promedia las distancias absolutas y ofrece una lectura directa del tamaño típico del error. R² compara el error del modelo con la variabilidad de los valores observados respecto de la media y puede ser negativo cuando el modelo predice peor que la media en ese conjunto. Chicco et al. (2021) discuten qué tan informativas son estas métricas y sus límites; aquí se informan juntas porque responden preguntas diferentes y el propósito es una comparación estudiantil sencilla, no declarar una métrica universal.')
p('Se evaluaron 15 000 filas de prueba para la regresión principal. Las métricas se calcularon una sola vez sobre la muestra guardada y se reportan redondeadas a tres decimales. Por la falta de una unidad explícita en el diccionario, MAE 5.165 y RMSE 8.144 se informan en la escala numérica de la variable. El documento no afirma todavía que sean microgramos por metro cúbico, aunque los boletines de SENAMHI expresan sus concentraciones en esa unidad; una definición general de una variable no sustituye una especificación de unidad en los metadatos del archivo.')
p('La evaluación gráfica complementa los promedios. En el diagrama real-estimado, una nube cercana a la diagonal indicaría que las estimaciones coinciden con los valores observados; la dispersión, en especial en los extremos, muestra errores que un solo número no resume. El boxplot y la serie mensual describen la distribución y su evolución, y no son mediciones de rendimiento fuera del test. La matriz de confusión y F1 pertenecen a la clasificación auxiliar: no deben mezclarse con las métricas de regresión ni presentarse como si midieran el mismo objetivo.')
p('El cuaderno calcula importancia por permutación con una submuestra de hasta 3 000 observaciones del test, tres repeticiones y score de −MAE. Se permuta una columna y se observa cuánto empeora ese score; un valor más alto señala que el estimador dependió más de esa información en esa evaluación. No mide un efecto causal, no separa variables correlacionadas y puede cambiar si cambia la muestra. Por eso se emplea como apoyo exploratorio y no como una explicación definitiva de la contaminación.')

head('21. Resultados',1)
head('Comparación principal',2)
p('Random Forest obtuvo el menor RMSE en la prueba temporal: MAE = 5.165, RMSE = 8.144 y R² = .499. Voting quedó segundo con RMSE = 8.300. El baseline de la media obtuvo RMSE = 11.885 y R² = −.067.')
table(6,'Métricas de regresión en la prueba temporal',['Modelo','MAE','RMSE','R²'],[
('Baseline de la media','9.325','11.885','−.067'),('Regresión lineal','6.748','9.604','.304'),('KNN regresor','6.823','9.774','.279'),('Árbol regresor','6.047','9.385','.335'),('Random Forest regresor','5.165','8.144','.499'),('Voting: lineal + Random Forest','5.505','8.300','.480')],[3.5,1,1,1],9)
figure(6,'RMSE y valores reales frente a estimados por Random Forest',models,'Métricas calculadas sobre 15 000 filas; dispersión de hasta 3 000 puntos, línea roja = igualdad entre real y estimado.')
head('Combinaciones y análisis complementarios',2)
p('Voting presentó correlación entre errores de .739. Su promedio no mejoró al Random Forest (RMSE 8.300 frente a 8.144). La mezcla ponderada eligió peso cero para lineal y uno para bosque en 2022; su resultado adicional fue MAE 6.486, RMSE 9.430, R² .328. Stacking de cuatro modelos y RidgeCV obtuvo MAE 6.064, RMSE 8.283, R² .482. Ninguna combinación superó al RF principal, pero la diferencia de periodo/volumen de entrenamiento limita una comparación causal definitiva.')
table(7,'Experimentos adicionales de combinación',['Experimento','MAE','RMSE','R²','Nota'],[
('RF principal','5.165','8.144','.499','Entrenamiento hasta 2022.'),('Mezcla ponderada','6.486','9.430','.328','0 % lineal + 100 % RF, base anterior a 2022.'),('Stacking + RidgeCV','6.064','8.283','.482','Bases anteriores a 2022; meta en 2022.')],[1.6,.65,.65,.5,3.1],8.2)
p('Clasificación y K-Means fueron complementarios. El umbral exploratorio de “PM2.5 relativamente alto” fue el percentil 75 del entrenamiento (31.70), no un límite sanitario. La proporción positiva fue .250 en train y .156 en test. Random Forest clasificador obtuvo F1 = .618 y Voting de logística + bosque F1 = .574. En K-Means, k = 2 tuvo el silhouette mayor entre k=2–6 (.425); las medianas de PM10/PM2.5/NO2 fueron 96.70/38.70/31.40 para el grupo 0 y 40.91/17.18/17.00 para el grupo 1. Los grupos no son categorías oficiales.')
p('La importancia por permutación fue mayor para PM10 (4.2881) y estación (1.4517). Esto expresa dependencia predictiva del modelo en esta muestra, no causalidad.')
figure(7,'Importancia por permutación del Random Forest',importance,'Cambio del score −MAE al permutar cada variable; valores mayores indican mayor dependencia predictiva en esa evaluación, no efecto causal.')

head('22. Discusión',1)
p('Random Forest superó al baseline y a los otros regresores en esta prueba. El promedio de árboles representó patrones de los datos con menor error que una regresión lineal. Aun así, R² = .499 indica ajuste parcial; no todos los cambios de PM2.5 se explican con las variables usadas. Peinado et al. (2025) revisan cientos de estudios y sostienen que la selección de modelos debe considerar datos disponibles, tipo de tarea y resolución, mientras que Karmoude et al. (2025) sintetizan retos que incluyen heterogeneidad espacial y temporal. Estas revisiones dan contexto al resultado, pero no permiten atribuir nuestro desempeño a un mecanismo concreto ni afirmar que un algoritmo sea mejor en todas las ciudades.')
p('El hallazgo sobre PM10 y la estación concuerda con el ranking de importancia por permutación del notebook, donde PM10 tuvo el mayor valor (4.2881) y luego estación (1.4517). Yang (2026) también presenta PM10 como predictor de interés en su estudio de Londres, pero su análisis es diario, usa otras entradas y formula una tarea de forecasting. La coincidencia entre dos resultados no prueba una causa común: en ambos casos el modelo puede apoyarse predictivamente en PM10, aunque el diseño y la escala sean diferentes. Además, la importancia por permutación puede repartir o distorsionar importancia cuando hay predictores correlacionados.')
p('Voting no mejoró al mejor componente. La correlación de errores .739 indica errores parcialmente parecidos, aunque no explica por sí sola la causa. Khadka et al. (2026) compararon Random Forest, SVR, LSTM y su promedio en series financieras, bajo particiones cronológicas y el mismo conjunto de características. En ese estudio el promedio quedó por debajo de LSTM y SVR, aunque por encima de Random Forest; su aplicación es bursátil y sus conclusiones no demuestran lo que ocurre con contaminantes. Aquí la explicación del resultado se basa primero en nuestras métricas: el promedio lineal–bosque obtuvo RMSE 8.300, frente a 8.144 del bosque individual. No puede concluirse que la regresión lineal “arrastre” siempre al Random Forest sin analizar errores caso por caso; la correlación de 0.739 sugiere errores relacionados y la ganancia efectiva debe decidirse midiendo el ensamble.')
p('Tian et al. (2024) reportaron un stacking para PM2.5 en Macao con LSTM, XGBoost y LightGBM. Singh et al. (2026) evaluaron un Voting ponderado para AQI de Taiwán, con cuatro modelos de boosting y explicación SHAP. Esos artículos muestran ensambles de calidad del aire con una estructura distinta a la nuestra; no contradicen que aquí el Voting no mejorara y tampoco hacen comparable el ranking de errores. El stacking de este proyecto quedó cerca del Random Forest principal, pero su entrenamiento terminó en 2021 y el meta-modelo se ajustó con 2022, mientras que el RF principal usó datos hasta 2022. Se requiere ejecutar ambos bajo un mismo protocolo temporal para concluir cuál combinación es más competitiva.')
p('Los trabajos locales de Montalvo et al. (2022), Rojas et al. (2022), Sánchez-Ccoyllo y Alonso (2024) y Gavidia et al. (2026) sitúan este ejercicio entre sistemas de bajo costo, simulación atmosférica y clasificación de contaminación en Lima. Muralikrishnan et al. (2026) estudian la predicción e interpretación de un índice AQI en un entorno industrial de India; su objetivo es una escala compuesta, distinto al valor de concentración que estima este notebook. Esta diferencia es relevante: la exactitud al predecir AQI no se convierte automáticamente en exactitud al estimar PM2.5.')
p('El aporte del proyecto es más limitado y verificable: comparar configuraciones vistas en clase sobre datos de SENAMHI, presentar gráficos y guardar en joblib el pipeline seleccionado para volver a usarlo con un esquema compatible. El modelo usa PM10 y NO2 contemporáneos. Si esos datos no están disponibles hasta después de medir PM2.5, el artefacto no puede funcionar como alerta anticipada. La interpretación correcta es estimación simultánea o un prototipo para reconstrucción experimental, no pronóstico futuro ni herramienta oficial de vigilancia.')

head('23. Amenazas a la validez y limitaciones',1)
for s in [
'Validez interna: la comparación principal usó muestras de 30 000 filas de entrenamiento y 15 000 de prueba, no todas las filas completas; otra muestra podría cambiar métricas.',
'Ensambles: sus experimentos usaron un periodo y tamaño de entrenamiento distinto al RF principal; se reportan como exploratorios.',
'Validez temporal: solo se evaluó un corte (2023–mayo de 2024); cambios futuros pueden alterar resultados.',
'Validez externa: siete estaciones del archivo no representan automáticamente toda Lima, Perú ni la red actual.',
'Validez de constructo: se estima PM2.5 contemporáneo con PM10 y NO2 de esa misma hora; no es pronóstico futuro ni alerta oficial.',
'Calidad: hay altos porcentajes de valores faltantes, especialmente NO2 (48.22 %); los casos completos podrían no representar toda la red.',
'Unidad: el diccionario descargado no especifica explícitamente unidades; confirmar con SENAMHI antes de atribuir unidad física a las métricas.',
'Validez de conclusión: no se calcularon intervalos de confianza ni se repitieron varios cortes temporales.'
]: bullet(s)
p('La primera limitación práctica es la muestra. El conjunto de modelado tenía 166 568 registros antes de 2023, pero la comparación principal tomó 30 000 aleatoriamente. Esto reduce el costo y mantiene la repetibilidad con semilla, pero puede dejar fuera combinaciones poco frecuentes de estación, temporada o concentraciones extremas. También se muestrearon 15 000 entre 42 534 filas de test; las métricas son una estimación para esa muestra y no tienen intervalos de confianza. Un siguiente estudio debería repetir el muestreo con varias semillas o evaluar todos los registros si el tiempo de cómputo lo permite.')
p('La segunda limitación corresponde al patrón de datos disponibles. Más de una tercera parte de PM2.5 y PM10 y casi la mitad de NO2 aparecen vacíos en la tabla original. Al filtrar casos completos se conservan registros utilizables sin inventar valores, pero las horas que cuentan con todos los sensores pueden ser distintas de las horas que no los tienen. Si la ausencia está concentrada en ciertas estaciones o periodos, los casos completos podrían representar de manera desigual el conjunto original. La limpieza no permite saber por qué faltan mediciones, y no se debe describir el filtro como imputación. Cicek y Aydin (2026) comparan estrategias de imputación para datos de contaminación en otro contexto; no se aplicó su método ni sus resultados al CSV de SENAMHI.')
p('La validez temporal se limita a un solo corte: antes de 2023 para ajustar los modelos principales y desde 2023 hasta mayo de 2024 para probarlos. El periodo posterior puede contener temporadas, eventos meteorológicos o cobertura distintos a otros años. Fauszt (2026) insiste en declarar qué información estaría disponible al momento de predecir; Mishra et al. (2026) y Nassar (2023) describen cómo el uso de información de validación o prueba puede producir resultados demasiado optimistas. El corte por fecha y la separación del preprocesamiento en pipeline reducen fugas convencionales, aunque aún se deben documentar el objetivo contemporáneo y el distinto protocolo de stacking.')
p('La validez de constructo depende del momento de entrada. Como el objetivo y los predictores incluyen mediciones de la misma hora, la prueba responde “¿qué PM2.5 estima el modelo si ya tengo PM10 y NO2 de esa hora?” y no “¿qué PM2.5 habrá mañana?”. Esa elección puede ser útil para un ejercicio de estimación simultánea o comparación, pero no permite presentar la salida como pronóstico de horas futuras. Para estudiar alertas sería necesario reformular la etiqueta como un valor futuro, utilizar solo información anterior al corte y volver a entrenar y evaluar el sistema.')
p('La validez externa es geográfica y temporal. El CSV utilizado contempla siete estaciones hasta mayo de 2024; no incluye por sí mismo toda la red de monitoreo actual ni estaciones de otras ciudades. El estudio de Quito de Alvarez et al. (2025) ilustra que la predicción de contaminantes puede depender de variables satelitales, topografía y cobertura espacial, aspectos ausentes del conjunto de características del notebook. Por tanto, resultados de estas siete estaciones no garantizan desempeño en otra estación, en otra región del Perú o en años siguientes con dinámicas climáticas diferentes.')
p('La validez de las comparaciones de ensamble está afectada por el protocolo. El Voting de la comparación principal comparte el test del RF principal, pero las pruebas de mezcla ponderada y stacking usan modelos base entrenados antes de 2022 y un conjunto de validación/meta de 2022. No tienen el mismo horizonte de entrenamiento del RF principal y el stacking usa predicciones de modelos prefijados para entrenar RidgeCV. Sus cifras son informativas para el parcial, aunque no aíslan de manera perfecta el efecto de combinar algoritmos. Una comparación más justa volvería a entrenar cada candidato con un mismo periodo y tamaño, ajustaría los pesos/meta-modelos solo en validación y preservaría un test final intacto.')
p('Por último, las métricas puntuales no cuantifican incertidumbre y el R² global puede ocultar errores diferentes entre estaciones o rangos de concentración. El valor de importancia por permutación depende de una muestra de hasta 3 000 filas y solo tres repeticiones. Para una futura versión convendría revisar MAE y sesgo por estación, revisar predicciones cerca de los extremos, repetir el cálculo de importancia con más repeticiones y reportar intervalos o resultados por varios cortes. Estas mejoras se presentan como recomendaciones y no como análisis ya ejecutados.')

head('24. Consideraciones éticas',1)
p('Se usaron registros públicos ambientales y no datos personales. Corresponde atribuir a SENAMHI y respetar la licencia indicada. La responsabilidad está en mantener la integridad de los datos, documentar faltantes y evitar comunicar una estimación como medición oficial (Alemohammad, 2026). El umbral P75 = 31.70 se eligió para practicar clasificación y no es ECA ni guía de la OMS.')
p('Ivanov (2025) propone principios de uso responsable de IA en investigación, entre ellos transparencia, trazabilidad, explicabilidad, supervisión humana y control de calidad. Aplicados de manera básica aquí, estos principios significan conservar el notebook y el `joblib`, registrar qué modelo se eligió, decir qué variables consume y exponer que los resultados solo aplican a los datos evaluados. También significa no presentar la variable “PM2.5 relativamente alto” del clasificador como categoría legal o sanitaria: el umbral es un percentil estadístico del entrenamiento.')
p('El conjunto puede ayudar a explicar mediciones, pero una salida del modelo no debe reemplazar el dato medido, el boletín del organismo ni la interpretación de especialistas. Este límite importa porque los errores son apreciables (R² = .499) y la tarea usa PM10 y NO2 simultáneos. Si se compartiera el prototipo, debería acompañarse de una nota visible sobre la tarea, periodo, estaciones y condición de entrada, además de respetar la licencia de atribución del dataset.')

head('25. Reproducibilidad',1)
p('El notebook registra Python 3.13.1, pandas 3.0.5, scikit-learn 1.9.1 y semilla 42. Para repetirlo, se descarga el CSV, se cambia la ruta en la configuración y se ejecutan las celdas en orden. Se guardan `resultados_senamhi/models/modelo_pm25_senamhi.joblib` y `resultados_senamhi/ficha_modelo.json`. El modelo requiere columnas compatibles; otra computadora debe ajustar la ruta y conservar el esquema del CSV.')

head('26. Conclusiones',1)
for s in [
'OE1: el CSV original contiene 577 794 filas, 15 columnas y siete estaciones entre enero de 2015 y mayo de 2024; PM2.5, PM10 y NO2 presentan valores faltantes.',
'OE2: se consolidaron 42 claves repetidas estación-hora (84 filas extra) y quedaron 577 710 registros. Para regresión se usaron 209 102 casos completos.',
'OE3/OE4: Random Forest tuvo el menor RMSE (8.144) y R² .499 en la muestra de prueba; superó al baseline, aunque deja variabilidad sin explicar.',
'Voting y las combinaciones adicionales no superaron al RF principal en los resultados reportados. El joblib guarda el pipeline RF individual; esto no demuestra que los ensambles sean siempre inferiores.',
'El modelo estima PM2.5 para la misma hora con PM10 y NO2 contemporáneos; no predice horas futuras ni sustituye una comunicación oficial.'
]: bullet(s)

head('27. Recomendaciones y trabajo futuro',1)
for s in [
'Confirmar con SENAMHI la unidad exacta de los campos del CSV.',
'Probar varios cortes temporales o validación de origen móvil.',
'Repetir Voting, mezcla y stacking con igual periodo y volumen de entrenamiento, reservando validación y test por separado.',
'Analizar errores por estación y por rango de PM2.5; añadir intervalos de incertidumbre.',
'Para pronosticar el futuro, usar variables disponibles antes del momento objetivo, como retardos horarios, y diseñar un nuevo test cronológico.',
'Si se hace una interfaz, presentarla como prototipo académico y aclarar que no sustituye mediciones ni alertas de SENAMHI.'
]: bullet(s)

doc.add_page_break(); head('28. Referencias',1)
refs=[
'Alemohammad, H. (2026). Data considerations for AI applications in environmental context. Field Actions Science Reports, Special Report, 34–39. https://journals.openedition.org/factsreports/8096',
'Alsabagh, A. S., Alawi, O. A., Kamar, H. M., Nafea, A. A., Al-Ani, M. M., Mohammed, H. A., Kazi, S. N., Oudah, A. Y., & Yaseen, Z. M. (2025). Deep learning framework for hourly air pollutants forecasting using encoding cyclical features across multiple monitoring sites in Beijing. Scientific Reports, 15, Article 22417. https://doi.org/10.1038/s41598-025-05472-5',
'Alvarez, C. I., Ulloa Vaca, C. A., & Echeverria Llumipanta, N. A. (2025). Machine learning for urban air quality prediction using Google AlphaEarth Foundations satellite embeddings: A case study of Quito, Ecuador. Remote Sensing, 17(20), Article 3472. https://doi.org/10.3390/rs17203472',
'Breiman, L. (2001). Random forests. Machine Learning, 45, 5–32. https://doi.org/10.1023/A:1010933404324',
'Chicco, D., Warrens, M. J., & Jurman, G. (2021). The coefficient of determination R-squared is more informative than SMAPE, MAE, MAPE, MSE and RMSE in regression analysis evaluation. PeerJ Computer Science, 7, Article e623. https://doi.org/10.7717/peerj-cs.623',
'Cicek, Z. I. E., & Aydin, Z. E. (2026). Enhancing air quality forecasting through missing data imputation: A stacking-based approach applied to urban monitoring data. PeerJ Computer Science, 12, Article e3904. https://doi.org/10.7717/peerj-cs.3904',
'Fang, Z., Yuan, Y., Zhang, J., Liu, Y., Mu, Y., Lu, Q., Xu, X., Wang, J., Wang, C., Zhang, S., & Chen, S. (2023). MLOps spanning whole machine learning life cycle: A survey. arXiv. https://doi.org/10.48550/arXiv.2304.07296',
'Fauszt, T. (2026). Time-consistent prediction in higher education: A framework for preventing data leakage in longitudinal models. Information, 17(6), Article 581. https://doi.org/10.3390/info17060581',
'Gavidia, A., Dominguez, A., & Flores-Chacón, E. (2026). Predicting air pollution in Metropolitan Lima using Gaussian Naïve Bayes (2025): An efficient model for urban environmental management. Sustainability, 18(11), Article 5748. https://doi.org/10.3390/su18115748',
'Ivanov, S. (2025). Responsible use of AI in social science research. The Service Industries Journal. https://doi.org/10.1080/02642069.2025.2537115',
'Izabi, M. B., Annas, S., & Ahmar, A. S. (2026). Evaluating Random Forest regression for air quality prediction. Jurnal Varian, 9(1), 77–84. https://doi.org/10.30812/varian.v9i1.6046',
'Karmoude, M., Munhungewarwa, B., Chiraira, I., Mckenzie, R., Kong, J., Smith, B., Ayana, G., Njara, N., Mathaha, T., Kumar, M., & Mellado, B. (2025). Machine learning for air quality prediction and data analysis: Review on recent advancements, challenges, and outlooks. Science of the Total Environment, 1002, Article 180593. https://doi.org/10.1016/j.scitotenv.2025.180593',
'Khadka, S., Thapa, P., Sharma, P., Silwal, S., & Kc, S. K. (2026). Do ensemble models always win? A comparative machine learning evaluation for financial time-series prediction. In Proceedings of the 2026 6th International Conference on Internet of Things and Machine Learning (IoTML 2026) (pp. 519–524). ACM. https://doi.org/10.1145/3838457.3838535',
'Kosmas, I., Papadopoulos, T., & Michalakelis, C. (2026). Machine learning lifecycle: A survey. AppliedMath, 6(7), Article 113. https://doi.org/10.3390/appliedmath6070113',
'Kovač, N., Ratković, K., Farahani, H., & Watson, P. (2024). A practical applications guide to machine learning regression models in psychology with Python. Methods in Psychology, 11, Article 100156. https://doi.org/10.1016/j.metip.2024.100156',
'Ministerio del Ambiente. (2017). Decreto Supremo N.° 003-2017-MINAM: Aprueban los estándares de calidad ambiental (ECA) para aire y establecen disposiciones complementarias. https://www.gob.pe/institucion/minam/normas-legales/3670-003-2017-minam',
'Mishra, L. N., Rangari, A., Nagrare, S., & Nayak, S. K. (2026). A taxonomy for detecting and preventing temporal data leakage in machine learning-based build prediction: A dual-platform empirical validation. PLoS ONE, 21(5), Article e0340167. https://doi.org/10.1371/journal.pone.0340167',
'Montalvo, L., Fosca, D., Paredes, D., Abarca, M., Saito, C., & Villanueva, E. (2022). An air quality monitoring and forecasting system for Lima City with low-cost sensors and artificial intelligence models. Frontiers in Sustainable Cities, 4, Article 849762. https://doi.org/10.3389/frsc.2022.849762',
'Mu, W., Cardelli, R., & Ferrari, S. (2026). Data preprocessing techniques for machine learning towards improving building energy performance: A systematic review. Energies, 19(6), Article 1561. https://doi.org/10.3390/en19061561',
'Muralikrishnan, R., Gollapalli, S., Sellappan, E., Prasanya, J., Murali, V., & Vijayakumar, A. (2026). Machine learning-driven prediction and interpretation of air quality index in industrial environment. Asian Journal of Civil Engineering, 27(3), 1473–1491. https://doi.org/10.1007/s42107-025-01572-9',
'Nassar, O. (2023). Data leakage in machine learning [Technical report]. https://doi.org/10.13140/RG.2.2.16250.81607',
'Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830. https://jmlr.org/papers/v12/pedregosa11a.html',
'Peffers, K., Rothenberger, M. A., Tuunanen, T., & Vaezi, R. (2012). Design science research evaluation. In K. Peffers, M. Rothenberger, & B. Kuechler (Eds.), Design science research in information systems: Advances in theory and practice (pp. 398–410). Springer. https://doi.org/10.1007/978-3-642-29863-9_29',
'Peinado, L. B., Guarda, T., Herrera-Vidal, G., Minnaard, C., & Coronado-Hernández, J. R. (2025). Statistical and machine learning models for air quality: A systematic review of methods and challenges. Algorithms, 18(12), Article 783. https://doi.org/10.3390/a18120783',
'Rojas, F. J., Pacsi-Valdivia, S., & Sánchez-Ccoyllo, O. (2022). Simulación computacional e influencia de las variables meteorológicas en las concentraciones de PM10 y PM2.5 en Lima Metropolitana. Información Tecnológica, 33(3), 223–238. https://doi.org/10.4067/S0718-07642022000300223',
'Sánchez-Ccoyllo, O. R., & Alonso, M. (2024). Improving PM10 and PM2.5 concentration prediction using the Brazilian Regional Atmospheric Modeling 5.2 System in Lima, Peru. Urban Climate, 55, Article 101985. https://doi.org/10.1016/j.uclim.2024.101985',
'Sarker, I. H. (2021). Machine learning: Algorithms, real-world applications and research directions. SN Computer Science, 2(3), Article 160. https://doi.org/10.1007/s42979-021-00592-x',
'SENAMHI. (2024a). Monitoreo de los contaminantes del aire en Lima Metropolitana [Conjunto de datos]. Plataforma Nacional de Datos Abiertos. https://www.datosabiertos.gob.pe/dataset/monitoreo-de-los-contaminantes-del-aire-en-lima-metropolitana-servicio-nacional-de',
'SENAMHI. (2024b). Vigilancia de la calidad del aire en el Área Metropolitana de Lima y Callao (AMLC): Enero 2024 [Boletín]. https://www.senamhi.gob.pe/load/file/03201SENA-130.pdf',
'SENAMHI. (2026). Vigilancia de la calidad del aire en el Área Metropolitana de Lima y Callao (AMLC): Agosto 2026 [Boletín]. https://www.senamhi.gob.pe/load/file/03201SENA-163.pdf',
'Singh, S., Kumar, M., Sengar, V., Kumar, A., Abhishek, K., & Shafeeq, B. M. A. (2026). Ensemble learning for air quality index prediction: Integrating gradient boosting, XGBoost, and stacking with SHAP-based interpretability. Scientific Reports, 16, Article 8544. https://doi.org/10.1038/s41598-026-39232-w',
'Solis Teran, M. A., Leite Coelho da Silva, F., Torres Armas, E. A., Carbo-Bustinza, N., & López-Gonzales, J. L. (2025). Modeling air pollution in Metropolitan Lima: A statistical and artificial neural network approach. Environments, 12(6), Article 196. https://doi.org/10.3390/environments12060196',
'Tian, H., Kong, H., & Wong, C. (2024). A novel stacking ensemble learning approach for predicting PM2.5 levels in dense urban environments using meteorological variables: A case study in Macau. Applied Sciences, 14(12), Article 5062. https://doi.org/10.3390/app14125062',
'Wirth, R., & Hipp, J. (2000). CRISP-DM: Towards a standard process model for data mining. In Proceedings of the 4th International Conference on the Practical Applications of Knowledge Discovery and Data Mining (pp. 29–39).',
'World Health Organization. (2021). WHO global air quality guidelines: Particulate matter (PM2.5 and PM10), ozone, nitrogen dioxide, sulfur dioxide and carbon monoxide. https://www.who.int/publications/i/item/9789240034228',
'Yang, X. (2026). Air quality forecasting in London using Random Forest regression. Journal of Innovation and Development, 14(2), 14–22. https://doi.org/10.54097/w7yyaw27',
'Yin, P.-Y. (2025). A review on PM2.5 sources, mass prediction, and association analysis: Research opportunities and challenges. Sustainability, 17(3), Article 1101. https://doi.org/10.3390/su17031101'
]
for ref in refs:
    x=doc.add_paragraph(); x.paragraph_format.left_indent=Inches(.5); x.paragraph_format.first_line_indent=Inches(-.5); x.paragraph_format.line_spacing=2; x.paragraph_format.space_after=Pt(0); x.paragraph_format.keep_together=True
    body_start=ref.find('). ')+3 if '). ' in ref else 0
    ranges=[]
    journals=['Field Actions Science Reports','Scientific Reports','Remote Sensing','Machine Learning','PeerJ Computer Science','Information','Sustainability','The Service Industries Journal','Jurnal Varian','Science of the Total Environment','AppliedMath','Methods in Psychology','PLoS ONE','Energies','Asian Journal of Civil Engineering','Journal of Machine Learning Research','Frontiers in Sustainable Cities','Información Tecnológica','Urban Climate','SN Computer Science','Algorithms','Environments','Applied Sciences','Journal of Innovation and Development']
    for journal in journals:
        at=ref.find(journal,body_start)
        if at>=0:
            ranges.append((at,at+len(journal)))
            m=re.match(r',\s*(\d+)',ref[at+len(journal):])
            if m: ranges.append((at+len(journal)+m.start(1),at+len(journal)+m.end(1)))
            break
    book_title='Design science research in information systems: Advances in theory and practice'
    at=ref.find(book_title)
    if at>=0: ranges.append((at,at+len(book_title)))
    ranges.sort()
    pos=0
    for a,b in ranges:
        if a>pos: font(x.add_run(ref[pos:a]),12)
        font(x.add_run(ref[a:b]),12,italic=True); pos=b
    if pos<len(ref): font(x.add_run(ref[pos:]),12)

head('29. Anexos',1)
head('Anexo A. Resultados auxiliares',2)
p('La clasificación y K-Means son complementarios y no cambian la pregunta principal de regresión.',False)
table(8,'Clasificación exploratoria de PM2.5 relativamente alto',['Modelo','Acc.','Prec.','Rec.','F1','ROC-AUC','PR-AUC'],[
('Baseline mayoritaria','.844','.000','.000','.000','—','—'),('Regresión logística','.839','.483','.465','.474','.839','.514'),('KNN clasificador','.843','.494','.317','.386','.765','.404'),('Árbol clasificador','.857','.542','.532','.537','.843','.522'),('Random Forest clasificador','.894','.710','.548','.618','.933','.678'),('Voting logística + RF','.888','.708','.483','.574','.920','.677')],[2.0,.65,.65,.65,.55,.85,.85],7.6)
p('El umbral fue el percentil 75 del entrenamiento (31.70); el porcentaje positivo cambió de 25.0 % en train a 15.6 % en test. No es un umbral sanitario.',False,size=9)
table(9,'Perfiles de K-Means con k = 2',['Grupo','Mediana PM10','Mediana PM2.5','Mediana NO2'],[('0','96.70','38.70','31.40'),('1','40.91','17.18','17.00')],[1.3,1.7,1.8,1.6],9)
p('Silhouette = .425, mayor entre k=2–6 probados. Las etiquetas de grupo son arbitrarias.',False,size=9)
head('Anexo B. Archivos del proyecto',2)
for s in ['Notebook: PARCIAL IA SENAHMI.ipynb.','Modelo: resultados_senamhi/models/modelo_pm25_senamhi.joblib.','Ficha: resultados_senamhi/ficha_modelo.json.','Dataset, metadatos y diccionario: Plataforma Nacional de Datos Abiertos / SENAMHI.']: bullet(s)

for t in doc.tables:
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs: par.paragraph_format.widow_control=True
doc.core_properties.title='Comparación de modelos de regresión para estimar PM2.5 en Lima Metropolitana'
doc.core_properties.subject='Paper académico con datos abiertos de SENAMHI'
doc.core_properties.author='Robert Eloy Herrera Ccari'
doc.core_properties.keywords='SENAMHI, PM2.5, Lima, aprendizaje automático, APA 7'
out=ROOT/'Paper_IA_Calidad_del_Aire_Lima_SENAMHI.docx'; doc.save(out)
print('CREATED',out,'ABSTRACT_WORDS',wc,'REFERENCES',len(refs),'TABLES',len(doc.tables),'PARAGRAPHS',len(doc.paragraphs))

